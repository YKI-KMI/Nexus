"""
NEXUS File Parser — Extract readable content from any file type.
Supports: PDF, DOCX, XLSX, code files, plaintext, and more.
"""

import os
import re
import chardet
from pathlib import Path

# Fix #6: module-level frozensets instead of functions that rebuild sets every call
CODE_EXTENSIONS = frozenset({
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".cpp", ".c", ".h",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".sh", ".bash",
    ".zsh", ".fish", ".ps1", ".lua", ".r", ".m", ".scala", ".clj",
    ".ex", ".exs", ".erl", ".hs", ".ml", ".fs", ".fsx", ".vue",
    ".graphql", ".proto", ".tf", ".hcl", ".bicep"
})

TEXT_EXTENSIONS = frozenset({
    ".txt", ".md", ".rst", ".log", ".json", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".env", ".xml", ".html",
    ".css", ".scss", ".sass", ".less", ".sql", ".gitignore",
    ".makefile", ".dockerfile", ".tex", ".bib", ".tsv", ".csv"
})


def extract_text(file_path: str, max_chars: int = 5000) -> str:
    """
    Universal text extractor. Returns text content or empty string.
    Never raises — always returns something.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".doc": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".xls": _extract_xlsx,
        ".csv": _extract_plaintext,
        ".pptx": _extract_pptx,
    }

    # Try extension-specific extractor
    if ext in extractors:
        try:
            text = extractors[ext](file_path)
            return text[:max_chars] if text else ""
        except Exception as e:
            pass

    # Code and text files
    if ext in CODE_EXTENSIONS or ext in TEXT_EXTENSIONS:
        try:
            return _extract_plaintext(file_path)[:max_chars]
        except Exception:
            pass

    # Try raw bytes + chardet for unknown files
    try:
        return _extract_raw(file_path)[:max_chars]
    except Exception:
        return ""


def get_file_metadata(file_path: str) -> dict:
    """Extract filename hints and basic metadata."""
    path = Path(file_path)
    name = path.stem.lower()
    
    # Clean up filename for hints
    hints = re.findall(r'[a-zA-Z]{3,}', name)
    
    return {
        "name": path.name,
        "stem": path.stem,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size if path.exists() else 0,
        "name_hints": " ".join(hints),
    }


def _extract_pdf(file_path: str) -> str:
    """Extract text from PDF using PyMuPDF."""
    import fitz  # PyMuPDF
    text_parts = []
    with fitz.open(file_path) as doc:
        for i, page in enumerate(doc):
            if i >= 10:  # Read first 10 pages max
                break
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _extract_docx(file_path: str) -> str:
    """Extract text from Word documents."""
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also grab table content
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _extract_xlsx(file_path: str) -> str:
    """Extract text from Excel files."""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.sheetnames[:3]:  # First 3 sheets
        ws = wb[sheet]
        parts.append(f"Sheet: {sheet}")
        rows_read = 0
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
                rows_read += 1
            if rows_read >= 50:
                break
    wb.close()
    return "\n".join(parts)


def _extract_pptx(file_path: str) -> str:
    """Extract text from PowerPoint files."""
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        parts = []
        for i, slide in enumerate(prs.slides):
            if i >= 10:
                break
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text)
        return "\n".join(parts)
    except ImportError:
        return ""


def _extract_plaintext(file_path: str) -> str:
    """Extract text from plaintext/code files with encoding detection.
    Fix #7: Try UTF-8 first (covers 95%+ of code/text files) and only
    fall back to chardet on decode failure to avoid the ~1-5 ms overhead."""
    # Fast path: try UTF-8 directly
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        pass

    # Slow path: detect encoding with chardet
    with open(file_path, "rb") as f:
        raw = f.read(8192)

    detected = chardet.detect(raw)
    encoding = detected.get("encoding") or "utf-8"

    try:
        with open(file_path, "r", encoding=encoding, errors="replace") as f:
            return f.read()
    except Exception:
        return raw.decode("utf-8", errors="replace")


def _extract_raw(file_path: str) -> str:
    """Last resort: read raw bytes and try to decode."""
    with open(file_path, "rb") as f:
        raw = f.read(4096)
    # Filter printable ASCII
    printable = bytes(b for b in raw if 32 <= b < 127 or b in (9, 10, 13))
    return printable.decode("ascii", errors="replace")


def is_binary_only(extension: str) -> bool:
    """Returns True for files where we can't easily extract text."""
    binary = {".exe", ".dll", ".so", ".dylib", ".bin", ".iso", ".img",
              ".zip", ".tar", ".gz", ".7z", ".rar", ".mp3", ".mp4",
              ".mov", ".avi", ".mkv", ".wav", ".flac", ".jpg", ".jpeg",
              ".png", ".gif", ".bmp", ".webp", ".svg", ".psd", ".ai"}
    return extension.lower() in binary


def describe_binary(file_path: str) -> str:
    """Generate description for binary files based on name and extension."""
    path = Path(file_path)
    name = path.stem
    ext = path.suffix.lower()

    type_map = {
        ".jpg": "image file", ".jpeg": "image file", ".png": "image file",
        ".gif": "animated image", ".webp": "image file", ".svg": "vector graphic",
        ".mp3": "audio file", ".wav": "audio file", ".flac": "audio file",
        ".mp4": "video file", ".mov": "video file", ".avi": "video file",
        ".zip": "compressed archive", ".tar": "archive", ".gz": "compressed file",
        ".psd": "photoshop design file", ".ai": "illustrator vector file",
        ".exe": "executable program", ".dll": "system library",
    }

    file_type = type_map.get(ext, f"{ext} file")
    
    # Try to extract meaning from filename
    words = re.findall(r'[a-zA-Z]{3,}', name)
    hint = " ".join(words[:6]) if words else name

    return f"{file_type}: {hint}"
